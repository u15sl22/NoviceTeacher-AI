from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import re
out=Path(r'C:\Users\13548\Desktop\NoviceTeacher-AI\output')
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.27); s.page_height=Inches(11.69)
s.top_margin=s.bottom_margin=Inches(.7)
s.left_margin=s.right_margin=Inches(.85)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name]; st.font.name='Times New Roman'; st.font.color.rgb=RGBColor(0,0,0)
normal=doc.styles['Normal']; normal.font.size=Pt(10.5)
normal.paragraph_format.line_spacing=1.0
normal.paragraph_format.space_after=Pt(5)
for name,size in [('Title',15),('Heading 2',12)]:
 st=doc.styles[name]; st.font.size=Pt(size); st.font.bold=True
 st.paragraph_format.space_before=Pt(10); st.paragraph_format.space_after=Pt(6)
 st.paragraph_format.keep_with_next=True
blocks=[
('title','3.X Guided Interaction Design'),
('p',"PedagoLoop structures the revision of existing lesson plans as a sequence of explicit teacher decisions. In an unrestricted dialogue, the connection between a proposed change, its instructional purpose, and the text it affects may become difficult to follow. A one-step rewrite may likewise combine several changes that a novice teacher must inspect together. These concerns motivate a design that makes individual issues and candidate revisions visible, without assuming that automation alone improves instructional planning. PedagoLoop divides the lesson plan into functional teaching sections and allows the teacher to review suggestions, request alternatives, and determine which changes enter the document."),
('p',"The interaction separates generating a candidate from applying it. AI output remains a proposal until the teacher explicitly accepts it. Revision therefore proceeds through manageable decisions about specific passages, with successive alternatives linked to the issue that prompted them. This organization is designed to support inspection and teacher control; its effects on revision quality and teacher development remain empirical questions."),
('h','3.X.1 Section-Based Revision Workflow'),
('p',"The teacher begins by pasting lesson-plan text or uploading a DOCX or PDF document. The resulting text is presented for confirmation before revision begins. This step allows the teacher to inspect the textual basis of subsequent suggestions, particularly when document extraction may affect how the original content is represented. The confirmed lesson plan is divided into functional teaching sections, and the system processes one section at a time rather than rewriting the complete document."),
('p',"Within the current section, PedagoLoop generates no more than two independent suggestions. Each identifies a specific issue, explains why a change is proposed, and supplies concrete revised text. Linking these components makes the proposed intervention inspectable: the teacher can consider the stated issue and the actual wording together. Section-level processing narrows attention to a local instructional problem, while the suggestion limit is intended to keep the immediate decision task manageable. The limit of two is a design choice, not an experimentally established optimum or evidence of reduced cognitive load."),
('p',"Each initial suggestion can develop into its own suggestion chain. A teaching section can be completed only when every active chain has received a final Accept or Reject decision. Requesting another candidate does not satisfy this condition. The completion rule makes unresolved proposals explicit and distinguishes the end of model generation from the end of the teacher's review."),
('h','3.X.2 Teacher Feedback Mechanism'),
('p',"The teacher responds to the current candidate through Accept, Reject, or Custom Revision. Accept approves that candidate and applies its proposed revision to the lesson plan. Reject ends consideration of the candidate without changing the existing text. Rejection records a decision not to use the proposal; it does not, by itself, establish why the teacher declined it or whether the underlying issue is unimportant."),
('p',"Custom Revision expresses a different position: the issue is worth addressing, but the proposed solution needs adjustment. The teacher provides a new instruction for improving the current candidate, and the system generates an alternative for review. This action avoids restricting feedback to a binary choice between adopting the AI's wording and abandoning the proposed change. It also preserves the distinction between requesting an alternative and approving its application. Table 1 summarizes the three actions."),
('table',''),
('h','3.X.3 Iterative Suggestion Chains'),
('p',"Repeated Custom Revision operations produce successive candidates for the same issue. These candidates form a suggestion chain rather than a set of unrelated suggestions. The recurring interaction is: Current suggestion → Teacher instruction → New candidate → Teacher decision. After examining the new candidate, the teacher may accept it, reject it, or issue another instruction. Custom Revision can be repeated until the latest candidate receives a final decision."),
('p',"Only the latest candidate in a chain is actionable. This rule establishes a single current proposal and prevents an earlier alternative from being treated as the active response after a later instruction has produced a new one. Previous candidates and teacher instructions remain in the interaction history, preserving the sequence through which the proposal changed. Their retention supports traceability without making every historical alternative simultaneously available for acceptance or rejection. Throughout the chain, the lesson-plan text remains unchanged unless the teacher explicitly accepts a candidate."),
('p',"As an illustrative example, suppose the AI identifies insufficient student practice and proposes adding a group activity. The teacher considers additional practice useful but requests a shorter individual activity because classroom time is limited. The system generates a new candidate addressing the same issue under that instruction. If the alternative is still too long, the teacher can request a further revision before deciding whether to accept it. This example illustrates the intended interaction, rather than an observed classroom outcome or experimental finding."),
('h','3.X.4 Context Continuity and Version Consistency'),
('p',"During repeated custom revision, the model receives the current relevant lesson-plan passage, the original issue, the previous candidate, and the teacher's latest instruction. Each component has a distinct role. The passage anchors generation in the document currently being revised; the original issue maintains the purpose of the chain; the previous candidate identifies the proposal being refined; and the latest instruction specifies the requested adjustment. Together, these inputs provide continuity for the immediate revision step without requiring the model to reconstruct the entire interaction."),
('p',"The system retains the original lesson plan, section versions, revision-round versions, AI suggestions, teacher decisions, custom instructions, and timestamps. This record supports tracing the relationship between generated alternatives, teacher actions, and document changes. However, the complete history is not inserted into every model request. Stored history serves traceability, whereas the selected request context serves the current generation task. This separation limits repeated inclusion of superseded material; it does not imply that every earlier instruction is automatically reconsidered at each step."),
('p',"Model generation is handled asynchronously so that a long generation request does not block the interface. Because the relevant document state may change before generation finishes, the system checks that the lesson-plan section and version are still current before saving the generated result. This check is designed to prevent a result produced for an older state from being carried forward as if it belonged to a newer version. Version consistency and explicit acceptance address separate requirements: a candidate must remain associated with the appropriate document state, and the teacher must authorize any resulting textual change."),
('h','3.X.5 Design Boundaries'),
('p',"The current mechanism records teacher feedback and supports iterative revision within suggestion chains. Custom instructions guide the generation of subsequent candidates, but the mechanism does not automatically learn a personalized policy or update model parameters. Retaining interaction records should therefore be understood as a means of documenting revision activity, rather than evidence of automatic personalization."),
('p',"PedagoLoop is designed to preserve teacher control by making proposals explicit, allowing alternatives to be requested, and requiring acceptance before modifying lesson-plan text. These are properties of the interaction mechanism, not established educational outcomes. Whether the mechanism improves lesson-plan quality or supports novice teachers' professional development requires empirical evaluation. The present design description makes no claim that these benefits have already been demonstrated.")]
texts=[]
for kind,txt in blocks:
 texts.append(txt)
 if kind=='title': doc.add_paragraph(txt,'Title')
 elif kind=='h': doc.add_paragraph(txt,'Heading 2')
 elif kind=='p': doc.add_paragraph(txt)
 else:
  p=doc.add_paragraph('Table 1. Teacher actions and their effects'); p.paragraph_format.keep_with_next=True
  p.paragraph_format.space_after=Pt(5)
  for r in p.runs: r.italic=True; r.font.size=Pt(10)
  rows=[['Teacher action','Meaning','System response','Effect on lesson-plan text'],['Accept','Approve the current candidate.','Apply the revision and close the chain.','Update the relevant text.'],['Reject','Decline the current candidate.','Record rejection and close the chain.','Retain existing text.'],['Custom Revision','Address the issue with a different solution.','Use the new instruction to generate another candidate in the chain.','No change pending explicit acceptance.']]
  t=doc.add_table(rows=0,cols=4); t.style='Table Grid'; t.autofit=False
  widths=[.85,1.55,2.05,1.95]
  for col,w in zip(t.columns,widths): col.width=Inches(w)
  for i,row in enumerate(rows):
   cells=t.add_row().cells
   for cell,w,txt in zip(cells,widths,row):
    cell.width=Inches(w); cell.text=txt; texts.append(txt)
    for p in cell.paragraphs:
     p.paragraph_format.keep_with_next=(i<3); p.paragraph_format.space_after=Pt(4); p.paragraph_format.space_before=Pt(4); p.paragraph_format.line_spacing=1.0
     for r in p.runs: r.font.size=Pt(10); r.bold=(i==0)
    if i==0:
     shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EFEFEF'); cell._tc.get_or_add_tcPr().append(shade)
   pr=t.rows[-1]._tr.get_or_add_trPr(); pr.append(OxmlElement('w:cantSplit'))
   if i==0: pr.append(OxmlElement('w:tblHeader'))
doc.core_properties.title='PedagoLoop Guided Interaction Design'
doc.core_properties.subject='Conference paper section on guided lesson-plan revision'
doc.core_properties.author=''
path=out/'PedagoLoop_Guided_Interaction_Design.docx'
for element in list(doc.styles.element.iter(qn('w:pBdr'))):
 element.getparent().remove(element)
for element in list(doc.element.iter(qn('w:pBdr'))):
 element.getparent().remove(element)
doc.save(path)
print(path)
print('Word count:',len(re.findall(r"\b[\w]+(?:[-’'][\w]+)*\b",' '.join(texts))))


